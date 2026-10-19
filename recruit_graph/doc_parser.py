# -*- coding: utf-8 -*-
"""把招聘场景里常见的文件（PDF / Word / 图片 / 文本）提取成纯文本。

供统一工作台 web_workbench.py 上传真实 JD、真实简历时使用：
- PDF：pymupdf 读文字层；文字很少的扫描页自动转成图片交给视觉模型 OCR；
- Word（.docx）：python-docx 读正文段落和表格；旧版 .doc 请先另存为 .docx；
- 图片（.png/.jpg/.jpeg/.webp/.bmp）：视觉模型 Qwen-VL 原样转录；
- 文本（.txt/.md）：直接按 UTF-8 读取。
"""

import base64
import os
from pathlib import Path

from dotenv import load_dotenv
from openai import OpenAI

load_dotenv(Path(__file__).resolve().parent.parent / ".env")

API_KEY = os.getenv("API_KEY")
BASE_URL = os.getenv("BASE_URL", "https://api.siliconflow.cn/v1")
VL_MODEL = os.getenv("VL_MODEL", "Qwen/Qwen3-VL-8B-Instruct")

_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
_MIME = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".bmp": "image/bmp",
    ".tif": "image/tiff",
    ".tiff": "image/tiff",
}


def _ocr_image_bytes(data: bytes, ext: str = ".png") -> str:
    """把一张图片的字节交给视觉模型，原样转录成文字。"""
    b64 = base64.b64encode(data).decode("utf-8")
    mime = _MIME.get(ext, "image/png")
    client = OpenAI(
        api_key=API_KEY,
        base_url=BASE_URL,
        timeout=120,
        max_retries=2,
    )
    resp = client.chat.completions.create(
        model=VL_MODEL,
        temperature=0,
        messages=[
            {
                "role": "user",
                "content": [
                    {
                        "type": "text",
                        "text": (
                            "这是一份简历或岗位描述的图片。请原样、逐字、完整地转录图中所有文字，"
                            "保留原有的段落、换行与条目顺序；只输出识别到的文字，"
                            "不要总结、不要翻译、不要补充、不要加任何说明。"
                        ),
                    },
                    {
                        "type": "image_url",
                        "image_url": {"url": f"data:{mime};base64,{b64}"},
                    },
                ],
            }
        ],
    )
    return resp.choices[0].message.content.strip()


def _extract_pdf(path: Path) -> str:
    """读 PDF：有文字层就直接取；近乎空白的扫描页转图片交给视觉模型。"""
    import pymupdf

    parts = []
    with pymupdf.open(path) as doc:
        for page in doc:
            text = page.get_text("text").strip()
            if len(text) >= 30:
                parts.append(text)
            else:
                # 扫描件：整页渲染成图，交给视觉模型
                pix = page.get_pixmap(dpi=200)
                parts.append(_ocr_image_bytes(pix.tobytes("png"), ".png"))
    return "\n".join(parts).strip()


def _extract_docx(path: Path) -> str:
    """读 Word 文档：正文段落 + 表格单元格（简历表格里常有信息）。"""
    import docx

    document = docx.Document(path)
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            line = " | ".join(cell for cell in cells if cell)
            if line:
                lines.append(line)
    return "\n".join(lines).strip()


def extract_text(file_path: str) -> str:
    """统一入口：按扩展名把文件提取成纯文本；不支持的类型抛清晰错误。"""
    path = Path(file_path)
    ext = path.suffix.lower()
    if not path.exists():
        raise FileNotFoundError(f"文件不存在：{path}")

    if ext == ".pdf":
        text = _extract_pdf(path)
    elif ext == ".docx":
        text = _extract_docx(path)
    elif ext == ".doc":
        raise ValueError("旧版 .doc 不支持，请用 Word/WPS 另存为 .docx 后再上传")
    elif ext in _IMAGE_EXTS:
        text = _ocr_image_bytes(path.read_bytes(), ext)
    elif ext in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8", errors="ignore").strip()
    else:
        raise ValueError(f"暂不支持的文件类型：{ext}（支持 PDF / Word / 图片 / 文本）")

    if not text.strip():
        raise ValueError("没能从文件里提取到文字，可能是空白或识别失败，请改粘文本")
    return text
